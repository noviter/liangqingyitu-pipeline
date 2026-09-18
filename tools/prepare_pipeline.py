#!/usr/bin/env python3
"""Prepare source materials for the Liangqingyitu review pipeline."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import fitz
import openpyxl
import xlrd
from docx import Document
from PIL import Image


# The manually maintained workbook contains high-resolution embedded images.
# We only read cell values, but openpyxl still touches image metadata on load.
Image.MAX_IMAGE_PIXELS = None


SUPPORTED_TEXT = {".pdf", ".docx"}
SOURCE_EXTENSIONS = {".xlsx", ".xls"}
SYSTEM_EXTENSIONS = {".pdf", ".doc", ".docx", ".zip", ".textclipping"}


def normalized(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def fingerprint(path: Path) -> str:
    digest = hashlib.sha1(str(path.relative_to(path.parents[1])).encode("utf-8")).hexdigest()
    return digest[:12]


def read_first_sheet(path: Path) -> list[list[Any]]:
    if path.suffix.lower() == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    return [sheet.row_values(row) for row in range(sheet.nrows)]


def value_after_label(rows: list[list[Any]], label: str, fallback_next_row: bool = False) -> str:
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if label in normalized(value):
                # The source forms use paired labels on the same row, e.g.
                # “单位 | ... | 部门 | ...”; select the adjacent value first.
                for candidate in row[column_index + 1 :]:
                    if normalized(candidate):
                        return normalized(candidate)
                if fallback_next_row and row_index + 1 < len(rows):
                    for candidate in rows[row_index + 1]:
                        if normalized(candidate):
                            return normalized(candidate)
    return ""


def responsibilities_from_source(path: Path) -> dict[str, Any]:
    rows = read_first_sheet(path)
    unit = value_after_label(rows, "单位")
    department = value_after_label(rows, "部门")
    responsibility = value_after_label(rows, "处室职能", fallback_next_row=True)
    if not department:
        department = path.stem

    return {
        "source_id": fingerprint(path),
        "source_file": str(path.as_posix()),
        "source_sheet": "first_sheet_only",
        "unit": unit,
        "department": department,
        "responsibility_text": responsibility,
        "status": "ready" if responsibility else "needs_review",
        "issue": "" if responsibility else "未识别到“处室职能”原文。",
    }


def extract_pdf(path: Path) -> tuple[str, str]:
    document = fitz.open(path)
    pages = []
    for number, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if text:
            pages.append(f"=== 第{number}页 ===\n{text}")
    return "\n\n".join(pages), f"pages={len(document)}"


def extract_docx(path: Path) -> tuple[str, str]:
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for table_number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [normalized(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            tables.append(f"=== 表{table_number} ===\n" + "\n".join(rows))
    return "\n\n".join(paragraphs + tables), f"paragraphs={len(paragraphs)}"


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def build_system_catalog(system_root: Path, output_root: Path, skip_text: bool) -> list[dict[str, Any]]:
    catalog: list[dict[str, Any]] = []
    text_root = output_root / "系统资料文本"
    for system_dir in sorted(path for path in system_root.iterdir() if path.is_dir()):
        match = re.match(r"\s*([0-9+]+)\s*(.*)", system_dir.name)
        system_code = match.group(1) if match else ""
        system_name = match.group(2).strip() if match else system_dir.name
        for document in sorted(path for path in system_dir.rglob("*") if path.is_file()):
            suffix = document.suffix.lower()
            item = {
                "evidence_id": fingerprint(document),
                "system_directory": system_dir.name,
                "system_code": system_code,
                "system_name": system_name,
                "source_file": str(document.as_posix()),
                "file_type": suffix or "[none]",
                "extract_status": "metadata_only",
                "extract_note": "未提取正文。",
                "text_file": "",
            }
            if suffix in SUPPORTED_TEXT and not skip_text:
                try:
                    text, note = extract_pdf(document) if suffix == ".pdf" else extract_docx(document)
                    if text:
                        text_path = text_root / f"{item['evidence_id']}.txt"
                        text_path.parent.mkdir(parents=True, exist_ok=True)
                        text_path.write_text(text, encoding="utf-8")
                        item.update({
                            "extract_status": "ready",
                            "extract_note": note,
                            "text_file": str(text_path.as_posix()),
                        })
                    else:
                        item.update({"extract_status": "empty_text", "extract_note": f"{note}; 未提取到可检索正文。"})
                except Exception as error:  # Keep a usable manifest even with one malformed file.
                    item.update({"extract_status": "failed", "extract_note": f"{type(error).__name__}: {error}"})
            elif suffix == ".doc":
                item.update({"extract_note": "旧版 DOC 未在首版自动提取；需要时转换为 DOCX/PDF 后重跑。"})
            elif suffix not in SYSTEM_EXTENSIONS:
                item.update({"extract_note": "非首版支持的资料格式，仅保留目录元数据。"})
            catalog.append(item)
    return catalog


def header_map(headers: list[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(headers):
        text = normalized(value)
        if "处室名称" in text:
            result["department"] = index
        elif text == "职责":
            result["responsibility"] = index
        elif "职能类型" in text:
            result["functional_type"] = index
        elif text == "事项":
            result["matter"] = index
        elif "数字化" in text:
            result["digital_status"] = index
        elif text.startswith("流程"):
            result["process"] = index
        elif "系统" in text:
            result["systems"] = index
        elif "依据" in text:
            result["evidence"] = index
        elif "待核验" in text:
            result["pending"] = index
    return result


def cell(row: tuple[Any, ...], mapping: dict[str, int], name: str) -> str:
    index = mapping.get(name)
    return normalized(row[index]) if index is not None and index < len(row) else ""


def read_manual_baseline(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    records = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        mapping = header_map(list(rows[0]))
        last_department = ""
        last_responsibility = ""
        sheet_records = []
        for row_number, row in enumerate(rows[1:], start=2):
            department = cell(row, mapping, "department") or last_department
            responsibility = cell(row, mapping, "responsibility") or last_responsibility
            matter = cell(row, mapping, "matter")
            if not matter:
                continue
            last_department = department
            last_responsibility = responsibility
            record = {
                "sheet": sheet.title,
                "row": row_number,
                "department": department,
                "responsibility": responsibility,
                "matter": matter,
                "functional_type": cell(row, mapping, "functional_type"),
                "digital_status": cell(row, mapping, "digital_status"),
                "process": cell(row, mapping, "process"),
                "systems": cell(row, mapping, "systems"),
                "evidence": cell(row, mapping, "evidence"),
                "pending": cell(row, mapping, "pending"),
            }
            sheet_records.append(record)
            records.append(record)
        sheets.append({"sheet": sheet.title, "header_map": mapping, "matter_count": len(sheet_records)})
    return {
        "source_file": str(path.as_posix()),
        "sheets": sheets,
        "records": records,
        "digital_distribution": Counter(record["digital_status"] or "空" for record in records),
    }


def drawio_pages(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    return re.findall(r'<diagram[^>]*\bname="([^"]+)"', text)


def report(sources: list[dict[str, Any]], catalog: list[dict[str, Any]], baseline: dict[str, Any], pages: list[str]) -> str:
    source_issues = [item for item in sources if item["status"] != "ready"]
    extraction = Counter(item["extract_status"] for item in catalog)
    lines = [
        "# 两清一图首版资料基线",
        "",
        f"生成时间：{datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')}",
        "",
        "## 职责原文",
        f"- 读取文件：{len(sources)}，均只读取第一个工作表。",
        f"- 已识别职责原文：{len(sources) - len(source_issues)}。",
        "",
        "## 系统资料",
        f"- 文档数量：{len(catalog)}。",
        f"- 正文可检索：{extraction['ready']}；无可检索正文：{extraction['empty_text']}；提取失败：{extraction['failed']}；仅目录元数据：{extraction['metadata_only']}。",
        "",
        "## 人工样例",
        f"- 处室页：{len(baseline['sheets'])}；事项记录：{len(baseline['records'])}。",
        "- 数字化结论分布：" + "；".join(f"{key}={value}" for key, value in baseline["digital_distribution"].items()) + "。",
        f"- 总泳道图页：{len(pages)}。",
        "",
        "## 进入后续 Skill 前的注意项",
        "- 将“metadata_only”“empty_text”“failed”的系统文件作为待补证据，而不是数字化结论。",
        "- 历史人工台账是对照基线，不覆盖新的原始职责或证据。",
        "- 事项、职责、证据和泳道图必须使用稳定 ID 关联，不按 Excel 行号关联。",
    ]
    if source_issues:
        lines.extend(["", "## 待处理职责来源"])
        lines.extend(f"- {item['source_file']}：{item['issue']}" for item in source_issues)
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare materials for the Liangqingyitu pipeline.")
    parser.add_argument("--source-dir", default="01_职责原文")
    parser.add_argument("--system-dir", default="02_系统资料")
    parser.add_argument("--sample-dir", default="04_人工梳理样例")
    parser.add_argument("--output-dir", default="工作区")
    parser.add_argument("--skip-text", action="store_true", help="Only index system documents; do not extract PDF/DOCX text.")
    args = parser.parse_args()

    source_dir = Path(args.source_dir)
    system_dir = Path(args.system_dir)
    sample_dir = Path(args.sample_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    sources = []
    for path in sorted(source_dir.iterdir()):
        if path.is_file() and path.suffix.lower() in SOURCE_EXTENSIONS:
            try:
                sources.append(responsibilities_from_source(path))
            except Exception as error:
                sources.append({
                    "source_id": fingerprint(path),
                    "source_file": str(path.as_posix()),
                    "source_sheet": "first_sheet_only",
                    "unit": "",
                    "department": path.stem,
                    "responsibility_text": "",
                    "status": "failed",
                    "issue": f"{type(error).__name__}: {error}",
                })

    catalog = build_system_catalog(system_dir, output_dir, args.skip_text)
    baseline = read_manual_baseline(sample_dir / "处室梳理表格.xlsx")
    pages = drawio_pages(sample_dir / "泳道图（总）.drawio")

    write_json(output_dir / "职责来源清单.json", sources)
    write_json(output_dir / "系统资料索引.json", catalog)
    write_json(output_dir / "人工样例基线.json", baseline)
    write_json(output_dir / "人工样例泳道图页签.json", pages)
    (output_dir / "资料基线报告.md").write_text(report(sources, catalog, baseline, pages), encoding="utf-8")
    print(f"Prepared {len(sources)} responsibility sources, {len(catalog)} system documents, {len(baseline['records'])} sample matters, and {len(pages)} drawio pages in {output_dir}.")


if __name__ == "__main__":
    main()
